"""Small helpers to add shapes to Artemiy's deck in the same visual style (Lato, navy 13316B, purple 5B3FA0)."""
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

NAVY, PURPLE, INK, MUTED, FAINT = "13316B", "5B3FA0", "1E2433", "5E6575", "7A8090"
CARD, WIN, LINE = "F1F3F8", "E6E0F5", "D5D9E3"
LEFT, WIDTH = 812800, 10566400


def rgb(h):
    return RGBColor.from_string(h)


def remove_placeholder_text(slide, startswith="[Filled in"):
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.startswith(startswith):
            sh._element.getparent().remove(sh._element)


def add_text(slide, x, y, w, h, text, size=14, bold=False, color=INK, align=None, anchor=None, fill=None, space_after=0, font="Lato"):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0 if fill is None else 127000)
    tf.margin_top = tf.margin_bottom = Emu(0 if fill is None else 63500)
    if anchor:
        tf.vertical_anchor = anchor
    if fill:
        tb.fill.solid()
        tb.fill.fore_color.rgb = rgb(fill)
    paras = text if isinstance(text, list) else [text]
    for i, ptxt in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        segs = ptxt if isinstance(ptxt, list) else [(ptxt, bold, color)]
        for seg in segs:
            t, b, c = (seg + (bold, color))[:3] if len(seg) < 3 else seg
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold = font, Pt(size), b
            r.font.color.rgb = rgb(c)
    return tb


def add_card(slide, x, y, w, h):
    s = slide.shapes.add_shape(1, Emu(x), Emu(y), Emu(w), Emu(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(CARD)
    s.line.fill.background()
    s.shadow.inherit = False
    bar = slide.shapes.add_shape(1, Emu(x), Emu(y), Emu(w), Emu(38100))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(NAVY)
    bar.line.fill.background()
    bar.shadow.inherit = False
    return s


def style_cell(cell, text, size=13, bold=False, color=INK, fill=None, align=PP_ALIGN.CENTER):
    cell.text = ""
    tf = cell.text_frame
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = "Lato", Pt(size), bold
    r.font.color.rgb = rgb(color)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Emu(60000)
    cell.margin_top = cell.margin_bottom = Emu(20000)
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
    else:
        cell.fill.background()


def add_chart(slide, x, y, w, h, categories, series: dict, kind="bar", colors=None, vmin=None, vmax=None,
              legend=True, font=12, number_format="0.00"):
    cd = CategoryChartData()
    cd.categories = categories
    for name, vals in series.items():
        cd.add_series(name, vals)
    ct = XL_CHART_TYPE.COLUMN_CLUSTERED if kind == "bar" else XL_CHART_TYPE.LINE_MARKERS
    gf = slide.shapes.add_chart(ct, Emu(x), Emu(y), Emu(w), Emu(h), cd)
    ch = gf.chart
    ch.font.size, ch.font.name = Pt(font), "Lato"
    ch.font.color.rgb = rgb(MUTED)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(font)
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("E4E6EC")
    va.format.line.fill.background()
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = number_format, False
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    ch.category_axis.format.line.color.rgb = rgb("B8BDC9")
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ch.category_axis.tick_labels.font.size = Pt(font)
    for i, s in enumerate(ch.plots[0].series):
        col = (colors or [NAVY, "EB6834", "1BAF7A", "EDA100", "8B6BD9"])[i % 5]
        if kind == "bar":
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = rgb(col)
        else:
            s.format.line.color.rgb = rgb(col)
            s.format.line.width = Pt(2.5)
            s.smooth = False
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = rgb(col)
            s.marker.format.line.color.rgb = rgb(col)
    return ch


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
