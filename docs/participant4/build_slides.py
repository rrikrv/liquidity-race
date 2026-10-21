from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]

ap = argparse.ArgumentParser()
ap.add_argument("--deck", required=True)
ap.add_argument("--out", default=str(ROOT / "slides" / "Do_Liquidity_Measures_Measure_Liquidity_p4.pptx"))
ap.add_argument("--tag", default="baseline")
args = ap.parse_args()

NAVY = RGBColor(19, 49, 107)
PURPLE = RGBColor(91, 63, 160)
INK = RGBColor(30, 36, 51)
GREY = RGBColor(94, 101, 117)
LIGHT = RGBColor(154, 161, 174)
CARD = RGBColor(0xF1, 0xF3, 0xF8)
GROUP = RGBColor(0xD9, 0xE0, 0xF2)
WHITE = RGBColor(255, 255, 255)
FONT = "Lato"

X0, WIDTH = 812800, 10566400

ORDER = ["roll_impact", "eff_tick_impact", "eff_tick2_impact", "holden_impact", "gibbs_impact",
         "lot_mixed_impact", "lot_ysplit_impact", "zeros_impact", "zeros2_impact", "amihud", "gamma", "amivest"]
HEAD = ["Roll Imp", "Eff Tick Imp", "Eff Tick2 Imp", "Holden Imp", "Gibbs Imp", "LOT Mix Imp",
        "LOT Y Imp", "Zeros Imp", "Zeros2 Imp", "Amihud", "Pas/Stam", "Amivest"]

PAPER = {
    ("A", "lambda"): [0.288, 0.296, 0.305, 0.305, 0.296, 0.309, 0.294, 0.278, 0.259, 0.317, -0.064, -0.030],
    ("A", "five"): [0.469, 0.439, 0.464, 0.463, 0.452, 0.467, 0.436, 0.384, 0.413, 0.516, 0.035, -0.139],
    ("B", "lambda"): [0.562, 0.404, 0.406, 0.505, 0.519, 0.482, 0.433, 0.485, 0.434, 0.400, -0.192, -0.062],
    ("B", "five"): [0.517, 0.293, 0.302, 0.397, 0.434, 0.394, 0.337, 0.491, 0.462, 0.511, -0.230, -0.170],
}
PAPER_BEST = {("A", "lambda"): "amihud", ("A", "five"): "amihud", ("B", "lambda"): "roll_impact", ("B", "five"): "roll_impact"}
PAPER_GROUP = {
    ("A", "lambda"): ["roll_impact", "eff_tick_impact", "eff_tick2_impact", "holden_impact", "gibbs_impact",
                      "lot_mixed_impact", "lot_ysplit_impact", "zeros_impact"],
    ("A", "five"): [],
    ("B", "lambda"): ["eff_tick_impact", "eff_tick2_impact", "holden_impact", "gibbs_impact", "lot_mixed_impact",
                      "lot_ysplit_impact", "zeros_impact", "zeros2_impact", "amihud"],
    ("B", "five"): ["holden_impact", "gibbs_impact", "lot_mixed_impact", "lot_ysplit_impact", "zeros_impact",
                    "zeros2_impact", "amihud"],
}
ROW_LABEL = {("A", "lambda"): "A  ·  Lambda", ("A", "five"): "A  ·  5-min impact",
             ("B", "lambda"): "B  ·  Lambda", ("B", "five"): "B  ·  5-min impact"}
OURS_BENCH = {"lambda": "lambda_true", "five": "post_impact_true"}


def drop(shape):
    el = shape._element
    el.getparent().remove(el)


def text(slide, x, y, w, h, s, size=14, bold=False, color=INK, align=None, anchor=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
    r.font.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


CARDS = [
    ("Amihud", "Amihud (2002)",
     "Price move per dollar traded: on an illiquid stock a small volume gives a big |r|",
     "Mean of |rₜ| / dollar volumeₜ over days with positive volume",
     "Best proxy of price impact: ρ = 0.52 with the 5-min impact, above all others; in the top group in all four races"),
    ("Extended Amihud", "new in this paper",
     "Drop the news term from Amihud: what is left is the spread per dollar traded",
     "Spread proxy ÷ average daily dollar volume. 9 versions (Roll Imp, Eff Tick Imp …); we compute 8, Holden is skipped",
     "Level with Amihud in three of the four races; Roll Imp is the best of the class"),
    ("Gamma", "Pastor, Stambaugh (2003)",
     "Order flow pushes the price and it reverts next day; γ < 0 is the reversal",
     "Monthly regression of rᵉₜ₊₁ on rₜ and sign(rᵉₜ)·Volumeₜ; γ is the slope",
     "Wrong sign or near zero: ρ = −0.06 with Lambda. Dominated"),
    ("Amivest", "liquidity ratio",
     "Dollars needed to move the price by 1%; larger means more liquid",
     "Mean of Volumeₜ / |rₜ| over days with non-zero return",
     "Negative by construction; small in size (−0.14 with the 5-min impact). Dominated"),
]


def slide_proxies(slide):
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.startswith("[Filled in"):
            drop(sh)
    gap = 190500
    cw = (WIDTH - 3 * gap) // 4
    y0, ch = 1234430, 4230000
    pad = 170000
    iw = cw - 2 * pad
    for i, (name, cite, model, est, res) in enumerate(CARDS):
        x = X0 + i * (cw + gap)
        rect(slide, x, y0, cw, ch, CARD)
        rect(slide, x, y0, cw, 38100, NAVY)
        text(slide, x + pad, y0 + 190000, iw, 330000, name, 20, True, NAVY)
        text(slide, x + pad, y0 + 540000, iw, 240000, cite, 12, False, GREY)
        for lab, body, ly, bh in (("MODEL", model, 880000, 900000), ("ESTIMATION", est, 1930000, 1150000),
                                  ("RESULT", res, 3130000, 950000)):
            text(slide, x + pad, y0 + ly, iw, 220000, lab, 11, True, PURPLE)
            text(slide, x + pad, y0 + ly + 215000, iw, bh, body, 12.5, False, INK)
    yb = y0 + ch + 190000
    rect(slide, X0, yb, 12700, 520000, PURPLE)
    tb = text(slide, X0 + 150000, yb, WIDTH - 150000, 520000,
              "Price impact is the harder target: correlations of 0.3–0.5 against up to 0.68 for the spread. "
              "Amihud is the one proxy built for it.", 13, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    notes(slide,
          "Четыре прокси price impact, всего около 30–40 секунд. "
          "Amihud: модуль доходности, делённый на долларовый объём за день, среднее по дням с объёмом. "
          "Расширенный Amihud: авторы раскладывают доходность на ликвидностную часть и новости, новости выбрасывают, "
          "остаётся спред на доллар объёма, поэтому берётся любой прокси спреда и делится на средний дневной долларовый объём; "
          "всего девять вариантов, мы считаем восемь, Holden пропущен вместе с Артемием. "
          "Gamma Пастора–Стамбо: регрессия доходности завтра на доходность сегодня и знаковый объём, "
          "коэффициент отрицательный, если после потока ордеров цена откатывается. "
          "Amivest: объём на единицу доходности, чем больше, тем ликвиднее, поэтому знак корреляции отрицательный. "
          "Главная мысль: Amihud единственный, кто сделан именно под price impact, и он выигрывает.")


def fmt(v):
    if v is None or pd.isna(v):
        return "—"
    return f"{v:.3f}".replace("-", "−")


def cell_set(cell, s, size=11, bold=False, color=INK, fill=None, align=PP_ALIGN.CENTER):
    cell.margin_left = cell.margin_right = Emu(30000)
    cell.margin_top = cell.margin_bottom = Emu(25000)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run()
    r.text = s
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
    r.font.color.rgb = color
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def load_ours(tag):
    wide = pd.read_csv(ROOT / "results" / f"table4_synthetic_{tag}.csv", index_col=0)
    lead = pd.read_csv(ROOT / "results" / f"table4_leaders_{tag}.csv")
    return wide, lead


def slide_races(slide, tag):
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.startswith("[Filled in"):
            drop(sh)
    wide, lead = load_ours(tag)
    keys = [("A", "lambda"), ("A", "five"), ("B", "lambda"), ("B", "five")]
    rows = 1 + 1 + 4 + 1 + 4
    cols = 13
    top = 1130000
    label_w = 1560000
    col_w = (WIDTH - label_w) // 12
    heights = [400000, 250000] + [262000] * 4 + [250000] + [262000] * 4
    gs = slide.shapes.add_table(rows, cols, Emu(X0), Emu(top), Emu(label_w + 12 * col_w), Emu(sum(heights)))
    tbl = gs.table
    tblPr = gs._element.graphic.graphicData.tbl.tblPr
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        from lxml import etree
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for a in ("firstRow", "bandRow"):
        tblPr.set(a, "0")
    tbl.columns[0].width = Emu(label_w)
    for j in range(1, 13):
        tbl.columns[j].width = Emu(col_w)
    for i, h in enumerate(heights):
        tbl.rows[i].height = Emu(h)

    cell_set(tbl.cell(0, 0), "", 10)
    for j, h in enumerate(HEAD):
        cell_set(tbl.cell(0, j + 1), h, 10, True, NAVY)

    def caption(r, s):
        tbl.cell(r, 0).merge(tbl.cell(r, 12))
        cell_set(tbl.cell(r, 0), s, 11, True, PURPLE, CARD, PP_ALIGN.LEFT)

    caption(1, "AUTHORS  ·  Table 4: TAQ 1993–2005, 400 stocks × 156 months")
    for k, key in enumerate(keys):
        r = 2 + k
        cell_set(tbl.cell(r, 0), ROW_LABEL[key], 11, True, INK, None, PP_ALIGN.LEFT)
        for j, c in enumerate(ORDER):
            v = PAPER[key][j]
            if c == PAPER_BEST[key]:
                cell_set(tbl.cell(r, j + 1), fmt(v), 11, True, WHITE, NAVY)
            elif c in PAPER_GROUP[key]:
                cell_set(tbl.cell(r, j + 1), fmt(v), 11, True, NAVY, GROUP)
            else:
                cell_set(tbl.cell(r, j + 1), fmt(v), 11, False, INK)
    meta = (ROOT / "results" / f"meta_{tag}.txt")
    label = meta.read_text().strip() if meta.exists() else "synthetic data"
    caption(6, f"OURS  ·  {label}")
    for k, key in enumerate(keys):
        r = 7 + k
        cell_set(tbl.cell(r, 0), ROW_LABEL[key], 11, True, INK, None, PP_ALIGN.LEFT)
        panel, bk = key
        bname = OURS_BENCH[bk]
        sub = lead[(lead["panel"] == panel) & (lead["benchmark"] == bname)].set_index("proxy")
        for j, c in enumerate(ORDER):
            if c not in sub.index:
                cell_set(tbl.cell(r, j + 1), "—", 11, False, LIGHT)
                continue
            v = sub.loc[c, "corr"]
            if bool(sub.loc[c, "best"]):
                cell_set(tbl.cell(r, j + 1), fmt(v), 11, True, WHITE, NAVY)
            elif bool(sub.loc[c, "in_leader_group"]):
                cell_set(tbl.cell(r, j + 1), fmt(v), 11, True, NAVY, GROUP)
            else:
                cell_set(tbl.cell(r, j + 1), fmt(v), 11, False, INK)

    yl = top + sum(heights) + 90000
    tb = slide.shapes.add_textbox(Emu(X0), Emu(yl), Emu(WIDTH), Emu(230000))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    for s, col in (("■ ", NAVY), ("best      ", GREY), ("■ ", RGBColor(0xB5, 0xC2, 0xE6)),
                   ("not significantly different from the best (5%)      ", GREY),
                   ("A: mean cross-sectional correlation   B: correlation of the equal-weighted portfolio over time", GREY)):
        r = p.add_run()
        r.text = s
        r.font.name, r.font.size = FONT, Pt(10.5)
        r.font.color.rgb = col
    return yl + 330000


def finish_races(slide, y, lines, note):
    rect(slide, X0, y, 12700, 700000, PURPLE)
    text(slide, X0 + 150000, y, WIDTH - 150000, 700000, lines, 13, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    notes(slide, note)


if __name__ == "__main__":
    prs = Presentation(args.deck)
    slide_proxies(prs.slides[5])
    s = prs.slides[9]
    y = slide_races(s, args.tag)
    takeaway = (ROOT / "slides" / f"takeaway_{args.tag}.txt")
    spoken = (ROOT / "slides" / f"notes_{args.tag}.txt")
    finish_races(s, y,
                 takeaway.read_text().strip() if takeaway.exists() else "TAKEAWAY",
                 spoken.read_text().strip() if spoken.exists() else "NOTES")
    prs.save(args.out)
    print("saved", args.out)
